from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
SOURCE_MD = WORKSPACE / "docs" / "prd" / "AI编导脚本生成工具_PRD.md"
OUTPUT_DOCX = WORKSPACE / "docs" / "prd" / "AI编导脚本生成工具_PRD.docx"

BODY_FONT = "Calibri"
EAST_ASIA_FONT = "Microsoft YaHei"
CONTENT_WIDTH_IN = 6.5
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGIN_DXA = 120


def set_run_font(run, *, size=None, bold=None, color=None, font=BODY_FONT):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph_spacing(paragraph, *, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, margin=CELL_MARGIN_DXA):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "bottom", "start", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(80 if side in ("top", "bottom") else margin))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color="666666")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color="666666")


def add_title(doc):
    title = doc.add_paragraph()
    set_paragraph_spacing(title, before=0, after=4, line=1.1)
    run = title.add_run("AI 编导脚本生成工具 MVP PRD")
    set_run_font(run, size=22, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    set_paragraph_spacing(subtitle, before=0, after=16, line=1.15)
    run = subtitle.add_run("短视频内容生产工作台产品需求方案")
    set_run_font(run, size=11, color="666666")


def configure_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
    normal.font.size = Pt(11)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), EAST_ASIA_FONT)
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("AI 编导脚本生成工具 MVP PRD")
    set_run_font(run, size=9, color="666666")

    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    add_title(doc)
    return doc


def split_table_row(line):
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def is_table_separator(line):
    cells = split_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def compute_widths(headers):
    col_count = len(headers)
    if col_count == 2:
        return [2400, CONTENT_WIDTH_DXA - 2400]
    if col_count == 3:
        return [2200, 3000, CONTENT_WIDTH_DXA - 5200]
    if col_count == 4:
        return [1900, 2500, 1800, CONTENT_WIDTH_DXA - 6200]
    if col_count == 5:
        return [1500, 1900, 1700, 2100, CONTENT_WIDTH_DXA - 7200]
    base = CONTENT_WIDTH_DXA // col_count
    widths = [base] * col_count
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    return widths


def add_markdown_table(doc, table_lines):
    rows = [split_table_row(line) for line in table_lines if not is_table_separator(line)]
    if not rows:
        return
    col_count = len(rows[0])
    rows = [row + [""] * (col_count - len(row)) for row in rows]
    widths = compute_widths(rows[0])

    table = doc.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    set_table_geometry(table, widths)

    for row_idx, row in enumerate(rows):
        for col_idx, value in enumerate(row[:col_count]):
            cell = table.cell(row_idx, col_idx)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_spacing(paragraph, before=0, after=0, line=1.1)
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, bold=(row_idx == 0), color="0B2545" if row_idx == 0 else "000000")
            if row_idx == 0:
                set_cell_shading(cell, "F2F4F7")

    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=6, line=1)


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [CONTENT_WIDTH_DXA])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    paragraph = cell.paragraphs[0]
    set_paragraph_spacing(paragraph, before=0, after=0, line=1.0)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, size=9, color="333333", font="Consolas")
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=6, line=1)


def add_paragraph_text(doc, text):
    paragraph = doc.add_paragraph()
    set_paragraph_spacing(paragraph, before=0, after=6, line=1.1)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def add_list_item(doc, text, ordered=False):
    paragraph = doc.add_paragraph(style="List Number" if ordered else "List Bullet")
    set_paragraph_spacing(paragraph, before=0, after=4, line=1.167)
    run = paragraph.add_run(text)
    set_run_font(run, size=11)


def build_docx():
    doc = configure_document()
    lines = SOURCE_MD.read_text(encoding="utf-8").splitlines()

    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()

        if line.startswith("# "):
            i += 1
            continue

        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                add_code_block(doc, code_lines)
                in_code = False
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue

        heading_match = re.match(r"^(#{2,4})\s+(.*)$", line)
        if heading_match:
            level = min(len(heading_match.group(1)) - 1, 3)
            paragraph = doc.add_heading(heading_match.group(2), level=level)
            for run in paragraph.runs:
                set_run_font(
                    run,
                    size={1: 16, 2: 13, 3: 12}[level],
                    bold=True,
                    color={1: "2E74B5", 2: "2E74B5", 3: "1F4D78"}[level],
                )
            i += 1
            continue

        ordered_match = re.match(r"^\d+\.\s+(.*)$", line)
        if ordered_match:
            add_list_item(doc, ordered_match.group(1), ordered=True)
            i += 1
            continue

        unordered_match = re.match(r"^[-*]\s+(.*)$", line)
        if unordered_match:
            add_list_item(doc, unordered_match.group(1), ordered=False)
            i += 1
            continue

        add_paragraph_text(doc, line)
        i += 1

    doc.save(OUTPUT_DOCX)


if __name__ == "__main__":
    build_docx()
    print(OUTPUT_DOCX)
